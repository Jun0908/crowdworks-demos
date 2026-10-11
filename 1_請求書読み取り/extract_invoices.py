#!/usr/bin/env python3
"""請求書PDF（テキストPDF / スキャンPDF）から支払情報を抽出し、支払期日順の Excel を出力するデモ。

使い方:
    python extract_invoices.py samples/ -o output/invoices.xlsx [--today 2026-10-11] [--self-name 自社名]

抽出項目: 請求日 / 合計金額(税込) / 取引先(発行元) / 振込先(銀行・支店・種別・番号・名義) / 支払期日 / 品目 → 勘定科目(ルール推定)
方式    : テキスト層があれば PyMuPDF で抽出、なければ tesseract(jpn) で OCR。項目は正規表現＋近傍ルール。
注意    : ルールベースのためレイアウトや表記の揺れに弱い。結果は必ず人が確認する前提（要確認フラグ参照）。
"""
from __future__ import annotations
import argparse, datetime as dt, json, re, sys, unicodedata
from dataclasses import dataclass, field, asdict
from pathlib import Path
from zoneinfo import ZoneInfo

import pymupdf
from PIL import Image, ImageOps
import numpy as np
import pytesseract
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEFAULT_SELF_NAME = "株式会社デモ受取ホールディングス"   # 自社名（宛名側を取引先と誤認しないため）
OCR_PSM = 6                                              # tesseract のページ分割モード（--ocr-psm で変更可）
MIN_TEXT_CHARS = 30                                      # これ未満のテキスト量ならスキャンとみなして OCR

# ----------------------------------------------------------------------------------------------
# テキスト取得
# ----------------------------------------------------------------------------------------------
CJK = r"\u3040-\u30ff\u3400-\u9fff\uff00-\uffef々〆ー"

def norm(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("\u3000", " ")
    # OCR や PDF で入る「日 本 語」の間の空白を詰める（ASCII 同士の空白は残す）
    s = re.sub(rf"(?<=[{CJK}])[ ](?=[{CJK}])", "", s)
    s = re.sub(rf"(?<=[{CJK}])[ ](?=[{CJK}])", "", s)
    return s

def text_lines_from_pdf(page) -> list[str]:
    """単語を y 座標で行にまとめ、大きな x ギャップは 2 スペース（列区切り）にする。"""
    words = page.get_text("words")  # x0,y0,x1,y1,word,block,line,wno
    words.sort(key=lambda w: (round((w[1] + w[3]) / 2), w[0]))
    rows: list[list] = []
    for w in words:
        yc = (w[1] + w[3]) / 2
        if rows and abs(rows[-1][0] - yc) < 3:
            rows[-1][1].append(w)
        else:
            rows.append([yc, [w]])
    lines = []
    for _, ws in rows:
        ws.sort(key=lambda w: w[0])
        out, prev = "", None
        for w in ws:
            if prev is not None:
                gap = w[0] - prev[2]
                out += "  " if gap > 15 else (" " if gap > 1.5 else "")
            out += w[4]; prev = w
        lines.append(out)
    return lines

def deskew(img: Image.Image) -> tuple[Image.Image, float]:
    """射影プロファイルで ±4° の傾きを推定して補正（簡易）。"""
    small = ImageOps.invert(img.convert("L")).resize((img.width // 4, img.height // 4))
    best, best_a = -1, 0.0
    for a in np.arange(-4, 4.01, 0.25):
        prof = np.asarray(small.rotate(a, resample=Image.BILINEAR)).astype(np.float32).sum(axis=1)
        score = float(np.var(prof))
        if score > best: best, best_a = score, a
    return img.rotate(best_a, resample=Image.BICUBIC, fillcolor=255, expand=False), float(best_a)

def ocr_lines(page) -> tuple[list[str], str]:
    pix = page.get_pixmap(dpi=300, colorspace=pymupdf.csGRAY)
    img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
    img, angle = deskew(img)
    img = ImageOps.autocontrast(img, cutoff=1)
    txt = pytesseract.image_to_string(img, lang="jpn", config=f"--psm {OCR_PSM} -c preserve_interword_spaces=1")
    return [l for l in txt.splitlines() if l.strip()], f"OCR(tesseract jpn psm{OCR_PSM}, deskew {angle:+.2f}°)"

def read_pdf(path: Path) -> tuple[list[str], str]:
    doc = pymupdf.open(path)
    lines, methods = [], set()
    for page in doc:
        if len(page.get_text().strip()) >= MIN_TEXT_CHARS:
            lines += text_lines_from_pdf(page); methods.add("text")
        else:
            ls, m = ocr_lines(page); lines += ls; methods.add(m)
    return [norm(l) for l in lines], " + ".join(sorted(methods))

# ----------------------------------------------------------------------------------------------
# 項目抽出
# ----------------------------------------------------------------------------------------------
DATE_PATTERNS = [
    re.compile(r"(?:令和|R)\s*(\d{1,2}|元)\s*[年.]\s*(\d{1,2})\s*[月.]\s*(\d{1,2})\s*日?"),
    re.compile(r"(20\d{2})\s*[年/\-.]\s*(\d{1,2})\s*[月/\-.]\s*(\d{1,2})\s*日?"),
]

def find_dates(s: str, with_end: bool = False):
    """(start, date) を出現順に返す（with_end=True なら (start, end, date)）。"""
    out = []
    for i, pat in enumerate(DATE_PATTERNS):
        for m in pat.finditer(s):
            try:
                if i == 0:
                    y = 2018 + (1 if m.group(1) == "元" else int(m.group(1)))
                else:
                    y = int(m.group(1))
                out.append((m.start(), m.end(), dt.date(y, int(m.group(2)), int(m.group(3)))))
            except ValueError:
                pass
    out.sort(key=lambda x: x[0])
    return out if with_end else [(a, d) for a, _, d in out]

ISSUE_LABEL = re.compile(r"(請求書?発行日|請求日|発行日|作成日|請求年月日|日付|Date)", re.I)
DUE_LABEL = re.compile(r"(お?支払い?(期日|期限|予定日|日)|お?振込み?(期日|期限)|ご?入金(期日|期限)|Due)", re.I)

def date_by_label(lines, label_re, lookahead=2):
    for i, l in enumerate(lines):
        m = label_re.search(l)
        if not m: continue
        ds = [d for p, d in find_dates(l) if p >= m.start()]
        if ds: return ds[0], f"ラベル「{m.group(0)}」"
        for j in range(1, lookahead + 1):
            if i + j < len(lines):
                ds = find_dates(lines[i + j])
                if ds: return ds[0][1], f"ラベル「{m.group(0)}」の{j}行下"
    return None, None

def extract_dates(lines):
    issue, issue_src = date_by_label(lines, ISSUE_LABEL)
    due, due_src = date_by_label(lines, DUE_LABEL)
    if due is None:  # 「◯年◯月◯日までに」型
        for l in lines:
            for p, e, d in find_dates(l, with_end=True):
                if re.match(r"\s*(まで|迄)", l[e:]):
                    due, due_src = d, "「…まで」表現"; break
            if due: break
    if issue is None:  # ラベルなし: 期日以外で最初に出てくる日付（多くは右上の日付）
        for l in lines[:15]:
            for _, d in find_dates(l):
                if d != due:
                    issue, issue_src = d, "ラベルなし(先頭付近の日付)"; break
            if issue: break
    return issue, issue_src, due, due_src

AMOUNT_LABELS = [  # 優先度順
    r"差引\s*ご?請求(金)?額",
    r"ご?請求(金)?額",
    r"お?支払(い)?(総|金)?額",
    r"(税込\s*合計|合計\s*金額|総合計|総額|合計\s*\(税込\))",
    r"(?<!小)合計",
]
NUM = re.compile(r"[¥\\]?\s*(\d{1,3}(?:[,.]\d{3})+(?![\d])|\d{4,})(?!\s*[%％])")

def to_int(s): return int(s.replace(",", "").replace(".", ""))

def amount_after(line, start):
    m = NUM.search(line, start)
    return to_int(m.group(1)) if m else None

def find_labeled_amount(lines, label):
    pat = re.compile(label)
    for i, l in enumerate(lines):
        m = pat.search(l)
        if m:
            v = amount_after(l, m.end())
            if v is None and i + 1 < len(lines):
                v = amount_after(lines[i + 1], 0)
            if v is not None: return v
    return None

def extract_amount(lines):
    total, src = None, None
    for lab in AMOUNT_LABELS:
        v = find_labeled_amount(lines, lab)
        if v is not None:
            total, src = v, lab; break
    sub = find_labeled_amount(lines, r"(小計|税抜(金額|合計)?)")
    tax = find_labeled_amount(lines, r"消費税(等)?")
    wh = find_labeled_amount(lines, r"源泉(徴収)?(税)?(額)?")
    check = None
    if total is not None and sub is not None and tax is not None:
        expect = sub + tax - (wh or 0)
        check = (expect == total)
    return total, src, check, (wh is not None)

COMPANY = re.compile(
    r"((?:株式会社|有限会社|合同会社|合資会社|合名会社|一般社団法人|税理士法人|\(株\)|\(有\))\s*[^\s:]+"
    r"|[^\s:]+?(?:株式会社|有限会社|合同会社|\(株\)|\(有\))"
    r"|[^\s:]+?(?:事務所|商店|商会|工房|スタジオ))")

def extract_vendor(lines, self_name):
    selfn = norm(self_name) if self_name else None
    cands = []
    reg_idx = [i for i, l in enumerate(lines) if re.search(r"T\d{13}", l.replace(" ", ""))]
    for i, l in enumerate(lines):
        for seg in re.split(r"\s{2,}", l):
            if re.search(r"(御中|様|宛|お客様|請求先)", seg): continue
            if re.search(r"(銀行|信用金庫|名義|振込|支店)", seg): continue
            for m in COMPANY.finditer(seg):
                name = m.group(1).strip()
                if selfn and (name in selfn or selfn in name): continue
                score = 0
                if any(abs(i - r) <= 3 for r in reg_idx): score += 3   # 登録番号の近く = 発行者の可能性大
                if i < len(lines) * 0.4: score += 1                    # 上部
                cands.append((score, -i, name))
    if not cands: return None, []
    cands.sort(reverse=True)
    return cands[0][2], [c[2] for c in cands]

BANK_RE = re.compile(r"([^\s:]+?(?:銀行|信用金庫|信用組合|労働金庫|信金|農業協同組合))")
BRANCH_RE = re.compile(r"([^\s:]+?(?:支店|営業部|出張所|本店))")
TYPE_RE = re.compile(r"(普通|当座|貯蓄)")
ACC_RE = re.compile(r"(?<![\d-])(\d{7})(?![\d-])")
KANA_HOLDER = re.compile(r"([ァ-ヶー][ァ-ヶー()・. ]*\)?[ァ-ヶー()・. ]*|[ァ-ヶ]\)[ァ-ヶー()・. ]+)")
LABEL_PREFIX = re.compile(r"^(【?お?振込先(口座)?】?|金融機関(名)?|銀行名|振込先)[:]?")

def extract_bank(lines):
    idx = next((i for i, l in enumerate(lines) if BANK_RE.search(l)), None)
    if idx is None: return {}
    win = lines[idx: idx + 6]
    text = "  ".join(win)
    r = {}
    m = BANK_RE.search(text)
    r["bank"] = LABEL_PREFIX.sub("", m.group(1)); pos = m.end()
    m = BRANCH_RE.search(text, pos)
    if m: r["branch"] = LABEL_PREFIX.sub("", m.group(1)); pos = m.end()
    m = TYPE_RE.search(text, pos)
    if m: r["type"] = m.group(1); pos = m.end()
    m = ACC_RE.search(text, pos) or ACC_RE.search(text)
    if m:
        r["number"] = m.group(1); after = m.end()
        # 名義: 「名義」ラベル優先、なければ口座番号直後のカナ
        hm = re.search(r"(?:口座)?名義(?:人)?\s*[:]?\s*([^\n]+?)(?:\s{2,}|$)", text)
        cand = hm.group(1) if hm else text[after:]
        km = KANA_HOLDER.search(cand.strip())
        if km: r["holder"] = re.sub(r"\s+", " ", km.group(1)).strip()
    return r

CATEGORY_KEYWORDS = {
    "通信費": ["インターネット", "回線", "通信", "電話", "携帯", "プロバイダ", "ドメイン", "SIM", "郵送", "切手", "宅配"],
    "消耗品費": ["コピー用紙", "用紙", "トナー", "インク", "文房具", "文具", "ボールペン", "付箋", "電池", "備品"],
    "外注費": ["業務委託", "委託", "外注", "制作", "デザイン", "開発", "代行", "コーディング", "翻訳"],
    "広告宣伝費": ["広告", "宣伝", "出稿", "掲載", "リスティング", "SNS", "PR", "チラシ", "ノベルティ"],
    "旅費交通費": ["交通費", "旅費", "出張", "宿泊", "航空券", "新幹線", "タクシー", "乗車券"],
    "地代家賃": ["賃料", "家賃", "共益費", "地代", "駐車場", "テナント"],
    "支払手数料": ["手数料", "仲介", "決済", "代引"],
    "水道光熱費": ["電気料", "電力", "ガス", "水道"],
}
ITEM_HEADER = re.compile(r"(品目|品名|摘要|内容|項目|明細|Description)")
ITEM_END = re.compile(r"(小計|合計|税抜|消費税)")

def extract_items(lines):
    start = next((i for i, l in enumerate(lines) if ITEM_HEADER.search(l)), None)
    if start is None: return []
    items = []
    for l in lines[start + 1:]:
        if ITEM_END.search(l): break
        name = re.split(r"\s{2,}|\s+\d[\d,]*(?:\s|$)|\s+1式", l)[0].strip()
        if name and not re.fullmatch(r"[\d,\s¥]+", name): items.append(name)
    return items

def estimate_category(items, all_lines):
    text = " ".join(items) if items else " ".join(all_lines)
    scores = {k: sum(text.count(w) for w in ws) for k, ws in CATEGORY_KEYWORDS.items()}
    ranked = sorted(scores.items(), key=lambda kv: -kv[1])
    if ranked[0][1] == 0: return "要確認", "キーワード該当なし"
    if len(ranked) > 1 and ranked[0][1] == ranked[1][1]:
        return "要確認", f"同点: {ranked[0][0]}/{ranked[1][0]}"
    detail = ", ".join(f"{k}:{v}" for k, v in ranked if v)
    return ranked[0][0], detail

# ----------------------------------------------------------------------------------------------
@dataclass
class Result:
    file: str
    method: str
    invoice_date: dt.date | None = None
    total: int | None = None
    vendor: str | None = None
    bank: dict = field(default_factory=dict)
    due_date: dt.date | None = None
    category: str | None = None
    items: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    needs_review: bool = False

def process(path: Path, self_name: str) -> Result:
    lines, method = read_pdf(path)
    r = Result(file=path.name, method=method)
    r.invoice_date, isrc, r.due_date, dsrc = extract_dates(lines)
    r.total, asrc, check, has_wh = extract_amount(lines)
    r.vendor, _ = extract_vendor(lines, self_name)
    r.bank = extract_bank(lines)
    r.items = extract_items(lines)
    r.category, cat_detail = estimate_category(r.items, lines)

    miss = [n for n, v in [("請求日", r.invoice_date), ("金額", r.total), ("取引先", r.vendor),
                           ("支払期日", r.due_date)] if v is None]
    bmiss = [k for k in ("bank", "branch", "type", "number", "holder") if not r.bank.get(k)]
    if miss: r.notes.append("未取得: " + "/".join(miss))
    if bmiss: r.notes.append("振込先の一部未取得: " + "/".join(bmiss))
    if check is False: r.notes.append("金額整合性NG(小計+税-源泉≠合計)")
    if check is None and r.total is not None: r.notes.append("金額の整合性チェック不可")
    if has_wh: r.notes.append("源泉徴収あり")
    if r.category == "要確認": r.notes.append(f"勘定科目要確認({cat_detail})")
    if isrc and "ラベルなし" in isrc: r.notes.append("請求日はラベルなし推定")
    if r.invoice_date and r.due_date and r.due_date < r.invoice_date: r.notes.append("期日<請求日")
    if "OCR" in method: r.notes.append("OCR読取: 目視確認推奨")
    r.needs_review = bool(miss or bmiss or check is False or r.category == "要確認" or "OCR" in method)
    r._debug = dict(lines=lines, issue_src=isrc, due_src=dsrc, amount_src=asrc, amount_check=check, category_detail=cat_detail)
    return r

def mark_duplicates(results):
    seen = {}
    for r in results:
        key = (r.vendor, r.total, r.invoice_date)
        if None in key: continue
        if key in seen:
            r.notes.append(f"重複の可能性: {seen[key]}"); r.needs_review = True
        else:
            seen[key] = r.file

# ----------------------------------------------------------------------------------------------
# Excel 出力
# ----------------------------------------------------------------------------------------------
HEADERS = [("支払期日", 12), ("残日数", 8), ("状態", 10), ("取引先", 28), ("支払金額(税込)", 15), ("請求日", 12),
           ("銀行", 20), ("支店", 14), ("種別", 6), ("口座番号", 10), ("口座名義", 30), ("勘定科目(推定)", 14),
           ("品目(抜粋)", 40), ("要確認", 8), ("確認理由・備考", 50), ("抽出方法", 26), ("ソースファイル", 32)]
FILL_OVERDUE = PatternFill("solid", fgColor="F8696B")
FILL_SOON = PatternFill("solid", fgColor="FFE699")
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_REVIEW = PatternFill("solid", fgColor="FCE4D6")

def write_excel(results, out: Path, today: dt.date):
    wb = Workbook(); ws = wb.active; ws.title = "支払予定一覧"
    ws.append([h for h, _ in HEADERS])
    thin = Side(style="thin", color="BFBFBF")
    for c, (_, w) in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    results = sorted(results, key=lambda r: (r.due_date is None, r.due_date or dt.date.max))
    for r in results:
        days = (r.due_date - today).days if r.due_date else None
        status = "期日不明" if days is None else ("期限超過" if days < 0 else ("7日以内" if days <= 7 else "—"))
        b = r.bank
        ws.append([r.due_date, days, status, r.vendor, r.total, r.invoice_date, b.get("bank"), b.get("branch"),
                   b.get("type"), b.get("number"), b.get("holder"), r.category, " / ".join(r.items[:4]),
                   "要確認" if r.needs_review else "", "; ".join(r.notes), r.method, r.file])
        row = ws.max_row
        fill = FILL_OVERDUE if (days is not None and days < 0) else (FILL_SOON if days is not None and days <= 7 else None)
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=c in (13, 15))
            if fill: cell.fill = fill
        if days is not None and days < 0:
            for c in range(1, 4): ws.cell(row=row, column=c).font = Font(bold=True, color="9C0006")
        if r.needs_review: ws.cell(row=row, column=14).fill = FILL_REVIEW if not fill else fill; ws.cell(row=row, column=14).font = Font(bold=True, color="C00000")
        ws.cell(row=row, column=1).number_format = "yyyy/mm/dd"
        ws.cell(row=row, column=6).number_format = "yyyy/mm/dd"
        ws.cell(row=row, column=5).number_format = '"¥"#,##0'
        ws.cell(row=row, column=10).number_format = "@"
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    n = ws.max_row
    ws.cell(row=n + 2, column=4, value="合計").font = Font(bold=True)
    tc = ws.cell(row=n + 2, column=5, value=f"=SUM(E2:E{n})"); tc.number_format = '"¥"#,##0'; tc.font = Font(bold=True)
    ws.cell(row=n + 3, column=1, value=f"基準日 {today:%Y/%m/%d}　赤=期限超過 / 黄=7日以内。ルールベース抽出のため、要確認行は必ず原本と照合してください。")
    # 生テキスト（検証用）
    ws2 = wb.create_sheet("抽出テキスト")
    ws2.append(["ソースファイル", "行", "テキスト"]); ws2.column_dimensions["A"].width = 32; ws2.column_dimensions["C"].width = 100
    for c in range(1, 4): ws2.cell(row=1, column=c).font = Font(bold=True)
    for r in results:
        for i, l in enumerate(r._debug["lines"], 1): ws2.append([r.file, i, l])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input_dir", type=Path)
    ap.add_argument("-o", "--output", type=Path, default=Path("output/invoices.xlsx"))
    ap.add_argument("--today", type=dt.date.fromisoformat, default=dt.datetime.now(ZoneInfo("Asia/Tokyo")).date())
    ap.add_argument("--self-name", default=DEFAULT_SELF_NAME, help="自社名（宛名として除外）")
    ap.add_argument("--json", type=Path, help="抽出結果を JSON でも保存")
    ap.add_argument("--ocr-psm", type=int, default=OCR_PSM, help="tesseract --psm (6=単一ブロック, 3=自動レイアウト解析)")
    a = ap.parse_args()
    globals()["OCR_PSM"] = a.ocr_psm
    pdfs = sorted(a.input_dir.glob("*.pdf"))
    if not pdfs: sys.exit(f"PDF が見つかりません: {a.input_dir}")
    results = []
    for p in pdfs:
        try:
            r = process(p, a.self_name)
        except Exception as e:  # 1件の失敗でバッチ全体を止めない
            r = Result(file=p.name, method="error", notes=[f"処理エラー: {e}"], needs_review=True); r._debug = dict(lines=[])
        results.append(r)
        print(f"[{r.method[:4]}] {p.name}: 期日={r.due_date} 金額={r.total} 取引先={r.vendor} 科目={r.category} {'※要確認' if r.needs_review else ''}")
    mark_duplicates(results)
    write_excel(results, a.output, a.today)
    if a.json:
        dump = []
        for r in results:
            d = asdict(r); d["debug"] = {k: v for k, v in r._debug.items() if k != "lines"}
            dump.append(d)
        a.json.write_text(json.dumps(dump, ensure_ascii=False, indent=2, default=str))
    print(f"→ {a.output} ({len(results)}件)")

if __name__ == "__main__":
    main()
