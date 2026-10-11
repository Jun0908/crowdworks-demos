import base64
import csv
import io
import json
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "3_契約更新抽出"))
sys.path.insert(0, str(ROOT / "web/dist/python"))
from browser_bridge import run_request

SAMPLE = (ROOT / "3_契約更新抽出/data/staff.csv").read_text(encoding="utf-8-sig")


def run(csv_text=SAMPLE, **kwargs):
    return json.loads(run_request({"csv_text": csv_text, **kwargs}))


def test_sample_and_all_downloads():
    r = run()
    assert (len(r["records"]), r["three_count"], len(r["missed_ids"]), len(r["issues"])) == (80, 37, 19, 7)
    wb = load_workbook(io.BytesIO(base64.b64decode(r["files"]["xlsx"])))
    assert "サマリー" in wb.sheetnames and "データ要確認" in wb.sheetnames
    rows = list(csv.DictReader(io.StringIO(base64.b64decode(r["files"]["csv"]).decode("utf-8-sig"))))
    assert len(rows) == 80
    with zipfile.ZipFile(io.BytesIO(base64.b64decode(r["files"]["zip"]))) as z:
        assert len(z.namelist()) == 13


def test_zero_results_export_and_independent_requests():
    zero = run(month="1900-01")
    assert zero["records"] == []
    load_workbook(io.BytesIO(base64.b64decode(zero["files"]["xlsx"])))
    assert len(run()["records"]) == 80


def test_lookahead_and_conventions_recompute():
    current = run()
    future = run(lookahead=1)
    assert len(future["records"]) > len(current["records"])
    for convention in ("minus1", "calendar"):
        r = run(convention=convention)
        assert r["convention"] == convention and r["files"]["xlsx"]


@pytest.mark.parametrize("payload", [
    {"csv_text": ""}, {"csv_text": "foo,bar\nx,y"}, {"csv_text": SAMPLE.splitlines()[0]},
    {"csv_text": SAMPLE, "month": "bad"}, {"csv_text": SAMPLE, "month": "1800-01"},
    {"csv_text": SAMPLE, "lookahead": 3}, {"csv_text": SAMPLE, "convention": "bad"},
])
def test_invalid_input_has_actionable_error(payload):
    with pytest.raises(ValueError):
        run_request(payload)


def test_csv_and_excel_formula_cells_are_text():
    text = '職員ID,氏名,契約開始日,契約期間(月),在籍状態,エリアマネージャー\nT1,=1+1,2026-08-01,3,在籍,=2+2\n'
    r = run(text, month="2026-10")
    exported = base64.b64decode(r["files"]["csv"]).decode("utf-8-sig")
    assert "'=1+1" in exported and "'=2+2" in exported
    wb = load_workbook(io.BytesIO(base64.b64decode(r["files"]["xlsx"])))
    assert not any(c.data_type == 'f' for ws in wb for row in ws for c in row)


def test_uploaded_markup_stays_plain_data():
    text = '職員ID,氏名,契約開始日,契約期間(月),在籍状態\nT1,<script>alert(1)</script>,2026-08-01,3,在籍\n'
    assert run(text, month="2026-10")["records"][0]["氏名"] == '<script>alert(1)</script>'
