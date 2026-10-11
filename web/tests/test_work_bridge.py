import base64
import copy
import io
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'web/dist/python'))
import invoice_core as I
from work_bridge import run_work

SAMPLE = json.loads((ROOT / 'web/dist/data/leads.json').read_text(encoding='utf-8'))
TEXT = ['請求日 2026年10月1日', '株式会社テスト通信', 'ご請求金額 11,000円',
        '支払期日 2026年10月15日', '品目 数量 金額', '通信回線 1式 10,000',
        '小計 10,000円', '消費税 1,000円', '架空銀行 さくら支店 普通 1234567',
        '口座名義 カ）テストツウシン']


def run(tool, **kwargs):
    return json.loads(run_work({'tool': tool, **kwargs}))


def invoice(**kwargs):
    return run('invoices', documents=[{'name': 'sample.pdf', 'lines': TEXT}], **kwargs)


def test_leads_matches_original_sample_results():
    result = run('leads', data=SAMPLE, sample=True)
    original = json.loads((ROOT / '2_焼肉店営業リスト/output/mock_leads_summary.json').read_text(encoding='utf-8'))
    for key in ['raw_results', 'after_place_id_dedupe', 'after_name_phone_address_dedupe',
                'final_rows', 'excluded', 'judgement', 'by_area']:
        assert result['summary'][key] == original[key]
    assert result['summary']['requests'] == {'searchText': 0, 'details': 0}
    workbook = load_workbook(io.BytesIO(base64.b64decode(result['files']['xlsx'])))
    assert workbook['リスト'].max_row == 15
    assert workbook['除外'].max_row == 5


def test_area_filter_recomputes_results():
    result = run('leads', data=SAMPLE, areas=['所沢'])
    assert len(result['records']) == 6
    assert len(result['excluded']) == 12
    assert {r['エリア'] for r in result['records']} == {'所沢'}


def test_missing_review_information_changes_evidence():
    result = run('leads', data=SAMPLE, use_reviews=False)
    assert result['summary']['judgement']['可能性'] == 0
    assert all('口コミに' not in r['判定根拠'] for r in result['records'])


@pytest.mark.parametrize('changes', [
    {'areas': []}, {'data': {'places': []}}, {'data': [{'id': '1'}]},
    {'data': [{'id': '1', 'displayName': '焼肉', 'formattedAddress': '東京都青梅市', 'reviews': ['bad']}]},
])
def test_invalid_leads_inputs(changes):
    with pytest.raises(ValueError):
        run('leads', **({'data': SAMPLE} | changes))


def test_uploaded_rows_render_as_data_and_export_safely():
    row = copy.deepcopy(SAMPLE['places'][0])
    row['displayName'] = {'text': '=焼肉<img src=x>'}
    row['websiteUri'] = 'javascript:alert(1)'
    result = run('leads', data=[row])
    assert result['records'][0]['店名'] == '=焼肉<img src=x>'
    workbook = load_workbook(io.BytesIO(base64.b64decode(result['files']['xlsx'])))
    assert workbook['リスト']['B2'].data_type != 'f'
    assert workbook['リスト']['B2'].value.startswith("'")
    assert workbook['リスト']['H2'].hyperlink is None


def test_leads_requests_do_not_mutate_input():
    before = copy.deepcopy(SAMPLE)
    run('leads', data=SAMPLE, use_reviews=False)
    assert SAMPLE == before
    assert run('leads', data=SAMPLE)['summary']['judgement']['可能性'] == 2


def test_invoice_fields_and_deadline_status():
    result = invoice(today='2026-10-11')
    row = result['records'][0]
    assert (row['total'], row['due_date'], row['invoice_date'], row['category']) == (11000, '2026-10-15', '2026-10-01', '通信費')
    assert row['bank']['number'] == '1234567'
    assert result['soon_count'] == 1
    assert invoice(today='2026-10-16')['overdue_count'] == 1


def test_ocr_spaces_in_thousands_separators_do_not_select_bank_number():
    result = run('invoices', documents=[{'name': 'scan.pdf', 'ocr': True,
                    'lines': ['差引ご請求額 199, 580', '振込先: 架空銀行 普通 7654321']}])
    assert result['records'][0]['total'] == 199580


def test_ocr_label_corrections_preserve_raw_evidence_and_review_flag():
    result = run('invoices', documents=[{'name': 'scan.pdf', 'ocr': True,
        'lines': ['株弐会社テスト', '支払期口 2026年10月16日', '合計金額 1,000円']}])
    row = result['records'][0]
    assert row['due_date'] == '2026-10-16'
    assert row['vendor'] == '株式会社テスト'
    assert row['lines'][0] == '株弐会社テスト'
    assert row['needs_review']
    assert any('OCR表記の補正' in note for note in row['notes'])


def test_invoice_ocr_and_missing_fields_are_flagged():
    result = run('invoices', documents=[{'name': 'scan.pdf', 'lines': TEXT, 'ocr': True},
                                      {'name': 'unknown.pdf', 'lines': ['ラベルのない書類']}])
    assert result['review_count'] == 2
    assert result['records'][-1]['total'] is None


def test_invoice_duplicate_detection_and_safe_raw_text_export():
    docs = [{'name': '=unsafe.pdf', 'lines': TEXT + ['=HYPERLINK("https://example.com")']},
            {'name': 'copy.pdf', 'lines': TEXT}]
    result = run('invoices', documents=docs)
    assert '重複の可能性' in ' '.join(result['records'][1]['notes'])
    workbook = load_workbook(io.BytesIO(base64.b64decode(result['files']['xlsx'])))
    assert workbook['支払予定一覧']['Q2'].value == "'=unsafe.pdf"
    assert all(c.data_type != 'f' for row in workbook['抽出テキスト'] for c in row)


@pytest.mark.parametrize('changes', [{'today': 'bad'}, {'documents': []}, {'self_name': 'x' * 201}])
def test_invalid_invoice_inputs(changes):
    with pytest.raises(ValueError):
        run('invoices', **({'documents': [{'name': 'a.pdf', 'lines': TEXT}]} | changes))
