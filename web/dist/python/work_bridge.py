"""Local, validated adapters for the invoice and sales-list browser demos."""
import base64
import copy
import csv
import io
import json
import tempfile
from collections import Counter
from dataclasses import asdict
from datetime import date
from pathlib import Path

import invoice_core as I
import leads_core as L


def safe(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    if isinstance(value, dict):
        return {k: safe(v) for k, v in value.items()}
    if isinstance(value, list):
        return [safe(v) for v in value]
    return value


def encoded(data):
    return base64.b64encode(data).decode()


def csv_bytes(columns, records):
    output = io.StringIO(newline='')
    writer = csv.DictWriter(output, fieldnames=columns, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(safe(records))
    return output.getvalue().encode('utf-8-sig')


def invoices(payload):
    documents = payload.get('documents')
    if not isinstance(documents, list) or not 1 <= len(documents) <= 5:
        raise ValueError('PDFは1〜5件を選択してください。')
    try:
        today = date.fromisoformat(payload.get('today', '2026-10-11'))
        if not 1900 <= today.year <= 2100:
            raise ValueError()
    except (ValueError, TypeError):
        raise ValueError('基準日は1900年〜2100年で指定してください。')
    self_name = payload.get('self_name', I.DEFAULT_SELF_NAME)
    if not isinstance(self_name, str) or len(self_name) > 200:
        raise ValueError('自社名は200文字以下にしてください。')
    results = []
    for doc in documents:
        if not isinstance(doc, dict) or not isinstance(doc.get('lines'), list):
            raise ValueError('読み取り結果の形式が不正です。')
        lines = doc['lines']
        if len(lines) > 10000 or any(not isinstance(l, str) or len(l) > 5000 for l in lines):
            raise ValueError('PDFの読み取りテキストが大きすぎます。')
        name = doc.get('name', 'invoice.pdf')
        if not isinstance(name, str) or len(name) > 200:
            raise ValueError('PDFのファイル名が不正です。')
        method = 'OCR (Tesseract.js / jpn)' if doc.get('ocr') else 'PDF text (PDF.js)'
        if doc.get('error'):
            result = I.Result(file=name, method='error', needs_review=True,
                              notes=['PDF読み取りエラー: ' + str(doc['error'])[:300]])
            result._debug = {'lines': []}
        else:
            result = I.process_lines(lines, name, method, self_name)
        results.append(result)
    I.mark_duplicates(results)
    results.sort(key=lambda r: (r.due_date is None, r.due_date or date.max))
    records = []
    for result in results:
        record = asdict(result)
        record['days'] = (result.due_date - today).days if result.due_date else None
        days = record['days']
        record['status'] = '期日不明' if days is None else '期限超過' if days < 0 else '7日以内' if days <= 7 else '予定'
        record['lines'] = result._debug.get('raw_lines', result._debug['lines'])
        records.append(record)
    export_rows = [{
        'ファイル': r['file'], '取引先': r['vendor'], '請求日': r['invoice_date'],
        '支払期日': r['due_date'], '金額': r['total'], '残日数': r['days'],
        '状態': r['status'], '勘定科目': r['category'], '要確認': '要確認' if r['needs_review'] else '',
        '銀行': r['bank'].get('bank'), '支店': r['bank'].get('branch'),
        '種別': r['bank'].get('type'), '口座番号': r['bank'].get('number'),
        '口座名義': r['bank'].get('holder'), '確認理由': ' / '.join(r['notes']),
    } for r in records]
    with tempfile.TemporaryDirectory() as directory:
        output = Path(directory) / 'invoices.xlsx'
        safe_results = copy.deepcopy(results)
        for r in safe_results:
            for key in ('file', 'vendor', 'bank', 'items', 'notes', 'category', '_debug'):
                setattr(r, key, safe(getattr(r, key)))
        I.write_excel(safe_results, output, today)
        files = {'xlsx': encoded(output.read_bytes()),
                 'csv': encoded(csv_bytes(list(export_rows[0]), export_rows)),
                 'json': encoded(json.dumps(records, ensure_ascii=False, indent=2, default=str).encode())}
    return {'records': records, 'today': today, 'files': files,
            'total': sum(r.total or 0 for r in results),
            'review_count': sum(r.needs_review for r in results),
            'overdue_count': sum(r['days'] is not None and r['days'] < 0 for r in records),
            'soon_count': sum(r['days'] is not None and 0 <= r['days'] <= 7 for r in records)}


def leads(payload):
    source = payload.get('data')
    if isinstance(source, dict):
        source = source.get('places')
    if not isinstance(source, list) or not 1 <= len(source) <= 1000:
        raise ValueError('検索結果JSONの places に1〜1,000件の店舗を入れてください。')
    if len(json.dumps(source, ensure_ascii=False).encode()) > 2 * 1024 * 1024:
        raise ValueError('JSONは2MB以下にしてください。')
    areas = payload.get('areas', list(L.AREA_BY_NAME))
    if not isinstance(areas, list) or not areas or any(a not in L.AREA_BY_NAME for a in areas):
        raise ValueError('対象エリアを1つ以上選んでください。')
    places = copy.deepcopy(source)
    for p in places:
        if not isinstance(p, dict) or not isinstance(p.get('id'), str) or not p['id'].strip():
            raise ValueError('各店舗に文字列の id が必要です。')
        for key in ('displayName', 'editorialSummary'):
            val = p.get(key, '')
            if isinstance(val, dict):
                val = val.get('text', '')
            if not isinstance(val, str) or len(val) > 5000:
                raise ValueError('店名・概要の形式が不正です。')
        if not L.place_text(p, 'displayName').strip() or not isinstance(p.get('formattedAddress'), str):
            raise ValueError('各店舗に店名 displayName と住所 formattedAddress が必要です。')
        for key in ('nationalPhoneNumber', 'businessStatus', 'primaryType', '_query', 'googleMapsUri', 'websiteUri'):
            if key in p and (not isinstance(p[key], str) or len(p[key]) > 1000):
                raise ValueError('店舗の文字列項目が不正です。')
        if not isinstance(p.get('types', []), list) or any(not isinstance(t, str) for t in p.get('types', [])):
            raise ValueError('types は文字列の配列にしてください。')
        if not isinstance(p.get('reviews', []), list) or any(not isinstance(r, dict) or not isinstance(r.get('text', {}), dict) or not isinstance(r.get('originalText', {}), dict) for r in p.get('reviews', [])):
            raise ValueError('口コミ reviews の形式が不正です。')
        for r in p.get('reviews', []):
            for key in ('text', 'originalText'):
                if not isinstance(r.get(key, {}).get('text', ''), str):
                    raise ValueError('口コミの本文は文字列にしてください。')
        p.pop('_queries', None)
        p.pop('_merged_ids', None)
        if not payload.get('use_reviews', True):
            p.pop('reviews', None)
            p.pop('editorialSummary', None)
    unique, stats = L.dedupe(places)
    rows, excluded = [], []
    note = '架空のサンプルデータ' if payload.get('sample') else 'ユーザー提供JSON / 外部検索は未実行'
    for p in unique:
        area = L.detect_area(p.get('formattedAddress'), L.AREAS)
        reason = '対象エリア外' if area is None or area.name not in areas else '閉業' if p.get('businessStatus') == 'CLOSED_PERMANENTLY' else '焼肉関連性なし' if not L.is_yakiniku(p) else None
        if reason:
            excluded.append({'除外理由': reason, '店名': L.place_text(p, 'displayName'),
                             '住所': p.get('formattedAddress'), '電話': p.get('nationalPhoneNumber', ''),
                             'place_id': p['id'], 'ヒット検索語': ' | '.join(p['_queries']), '備考': note})
        else:
            rows.append(L.to_row(p, area, note))
    order = {a.name: i for i, a in enumerate(L.AREAS)}
    rank = {'確定': 0, '可能性': 1, '不明': 2}
    rows.sort(key=lambda r: (order[r['エリア']], rank[r['炭火判定']], r['店名']))
    summary = {'mode': 'mock' if payload.get('sample') else 'uploaded-json',
               'requests': {'searchText': 0, 'details': 0}, 'raw_results': stats['input'],
               'after_place_id_dedupe': stats['after_place_id'],
               'after_name_phone_address_dedupe': stats['after_name_phone_address'],
               'final_rows': len(rows), 'excluded': dict(Counter(r['除外理由'] for r in excluded)),
               'judgement': {k: sum(r['炭火判定'] == k for r in rows) for k in rank},
               'by_area': {a: sum(r['エリア'] == a for r in rows) for a in areas}}
    with tempfile.TemporaryDirectory() as directory:
        output = Path(directory) / 'leads.xlsx'
        L.write_xlsx(output, safe(rows), safe(excluded), summary, bool(payload.get('sample')))
        from openpyxl import load_workbook
        workbook = load_workbook(output)
        for sheet in workbook:
            for row in sheet:
                for c in row:
                    if c.hyperlink and not c.hyperlink.target.startswith(('https://', 'http://')):
                        c.hyperlink = None
        workbook['注記'].cell(3, 1, 'データ出典: 同梱の架空サンプルまたはユーザー提供JSON。ブラウザ版はGoogle Places APIへアクセスしません。')
        workbook.save(output)
        files = {'xlsx': encoded(output.read_bytes()),
                 'csv': encoded(csv_bytes(L.OUTPUT_COLUMNS, rows)),
                 'excluded_csv': encoded(csv_bytes(L.EXCLUDED_COLUMNS, excluded))}
    return {'records': rows, 'excluded': excluded, 'summary': summary, 'files': files}


def run_work(payload):
    tool = payload.get('tool')
    if tool not in ('invoices', 'leads'):
        raise ValueError('処理の種類が不正です。')
    output = invoices(payload) if tool == 'invoices' else leads(payload)
    return json.dumps(output, ensure_ascii=False, default=str)
