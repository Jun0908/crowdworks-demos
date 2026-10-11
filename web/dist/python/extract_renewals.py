"""契約更新対象者抽出 CLI（デモ）

例:
  python extract_renewals.py --month 2026-11
  python extract_renewals.py --month 2026-11 --lookahead 1          # 翌月末まで
  python extract_renewals.py --month 2026-11 --convention minus1    # 終了日規約の切替
  python extract_renewals.py --month 2026-11 --deadline 2026-10-20

出力（--out 配下、既定 output/）:
  更新対象者_YYYY-MM.xlsx   サマリー / AM別シート / 除外 / データ要確認 / 旧方式比較
  更新対象者_YYYY-MM.csv    対象者一覧（UTF-8 BOM）
  messages_YYYY-MM/*.txt    AM別 確認依頼文テンプレート（送信はしない）
※ 本スクリプトはメール等の送信を一切行わない。
"""
from __future__ import annotations

import argparse
import csv
import os
import re
from collections import Counter, defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

import renewal as R

WD = "月火水木金土日"


def jdate(d: date) -> str:
    return f"{d.year}年{d.month}月{d.day}日（{WD[d.weekday()]}）"


def load_rows(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def target_record(t: R.Target) -> dict:
    r = t.row
    return {
        "職員ID": r[R.COL_ID], "氏名": r[R.COL_NAME], "エリア": r.get(R.COL_AREA, ""),
        "エリアマネージャー": (r.get(R.COL_AM) or "").strip() or R.AM_MISSING_LABEL,
        "契約期間(月)": t.months, "当初契約開始日": t.anchor, "現契約開始日": t.start,
        "現契約終了日(計算)": t.end, "第何期": t.k, "更新回数": t.renewals_so_far,
        "契約終了日列(元データ)": r.get(R.COL_END, ""),
        "雇止め予告期限(30日前)": t.notice_deadline if t.notice_required else "",
        "無期転換見込": "要確認" if t.mukitenkan else "",
        "フラグ": " / ".join(t.flags), "更新可否(AM記入)": "", "備考(AM記入)": "",
    }


# ------------------------------------------------------------------ xlsx
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
THREE_FILL = PatternFill("solid", fgColor="E2EFDA")
INPUT_FILL = PatternFill("solid", fgColor="FFFFE0")


def _autowidth(ws, min_w=8, max_w=48, from_row=1):
    widths = defaultdict(int)
    for row in ws.iter_rows(min_row=from_row):
        for c in row:
            if c.value is None:
                continue
            v = c.value.isoformat() if isinstance(c.value, date) else str(c.value)
            w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in v)
            widths[c.column] = max(widths[c.column], w)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w + 2))


def _write_table(ws, start_row, headers, records, name=None, highlight=True):
    for j, h in enumerate(headers, 1):
        c = ws.cell(start_row, j, h)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, rec in enumerate(records, start_row + 1):
        for j, h in enumerate(headers, 1):
            v = rec.get(h, "")
            c = ws.cell(i, j, v)
            c.border = BORDER
            if isinstance(v, date):
                c.number_format = "yyyy/mm/dd"
                c.alignment = Alignment(horizontal="center")
        if highlight:
            if rec.get("フラグ"):
                ws.cell(i, headers.index("フラグ") + 1).fill = FLAG_FILL
            if rec.get("契約期間(月)") == 3:
                ws.cell(i, headers.index("契約期間(月)") + 1).fill = THREE_FILL
            for h in ("更新可否(AM記入)", "備考(AM記入)"):
                if h in headers:
                    ws.cell(i, headers.index(h) + 1).fill = INPUT_FILL
    if name and records:
        ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(records)}"
        tab = Table(displayName=name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tab)


def _safe_sheet(name, used):
    n = re.sub(r"[\[\]\*\?/\\:]", "_", name)[:28]
    base, i = n, 2
    while n in used:
        n = f"{base[:25]}_{i}"; i += 1
    used.add(n)
    return n


def build_xlsx(path, res: R.ExtractResult, recs, naive, deadline, period_label):
    wb = Workbook()
    ws = wb.active
    ws.title = "サマリー"
    ws["A1"] = f"契約更新対象者 サマリー（{period_label}満了分）"; ws["A1"].font = TITLE_FONT
    ws["A2"] = f"終了日規約: {res.convention} ／ 確認回答期限: {jdate(deadline)} ／ 作成日: {date.today().isoformat()} ／ ※架空データによるデモ"
    ws["A3"] = f"対象者 {len(recs)} 名 ／ 退職予定のため除外 {len(res.excluded)} 名 ／ データ要確認 {len(res.issues)} 件"
    ws["A3"].font = Font(bold=True)
    row = 5
    by_am = Counter((r["エリア"], r["エリアマネージャー"]) for r in recs)
    by_am_len = defaultdict(Counter)
    for r in recs:
        by_am_len[(r["エリア"], r["エリアマネージャー"])][r["契約期間(月)"]] += 1
    ws.cell(row, 1, "■ エリアマネージャー別").font = Font(bold=True); row += 1
    hdr = ["エリア", "エリアマネージャー", "対象者数", "1ヶ月", "3ヶ月", "6ヶ月", "12ヶ月", "フラグ付き"]
    flagged = Counter((r["エリア"], r["エリアマネージャー"]) for r in recs if r["フラグ"])
    data = [{"エリア": a, "エリアマネージャー": m, "対象者数": n,
             "1ヶ月": by_am_len[(a, m)][1], "3ヶ月": by_am_len[(a, m)][3],
             "6ヶ月": by_am_len[(a, m)][6], "12ヶ月": by_am_len[(a, m)][12],
             "フラグ付き": flagged[(a, m)]} for (a, m), n in sorted(by_am.items())]
    data.append({"エリア": "合計", "エリアマネージャー": "", "対象者数": len(recs),
                 **{f"{k}ヶ月": sum(d[f"{k}ヶ月"] for d in data) for k in (1, 3, 6, 12)},
                 "フラグ付き": sum(flagged.values())})
    _write_table(ws, row, hdr, data, highlight=False)
    for j in range(1, len(hdr) + 1):
        ws.cell(row + len(data), j).font = Font(bold=True)
    row += len(data) + 2

    ws.cell(row, 1, "■ エリア別").font = Font(bold=True); row += 1
    by_area = Counter(r["エリア"] for r in recs)
    _write_table(ws, row, ["エリア", "対象者数"], [{"エリア": a, "対象者数": n} for a, n in sorted(by_area.items())], highlight=False)
    row += len(by_area) + 2
    ws.cell(row, 1, "■ 契約期間別").font = Font(bold=True); row += 1
    by_len = Counter(r["契約期間(月)"] for r in recs)
    _write_table(ws, row, ["契約期間(月)", "対象者数"], [{"契約期間(月)": k, "対象者数": by_len[k]} for k in (1, 3, 6, 12)], highlight=False)
    row += 6
    ws.cell(row, 1, "■ 旧方式との比較（3ヶ月契約者）").font = Font(bold=True); row += 1
    _write_table(ws, row, ["方式", "抽出数(3ヶ月)", "正しい対象(3ヶ月)", "漏れ", "誤抽出"], naive, highlight=False)
    _autowidth(ws)
    ws.column_dimensions["A"].width = 36
    ws.freeze_panes = "A5"
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["B"].width = 24
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

    headers = ["職員ID", "氏名", "エリア", "契約期間(月)", "当初契約開始日", "現契約開始日",
               "現契約終了日(計算)", "第何期", "更新回数", "契約終了日列(元データ)",
               "雇止め予告期限(30日前)", "無期転換見込", "フラグ", "更新可否(AM記入)", "備考(AM記入)"]
    used = {"サマリー"}
    groups = defaultdict(list)
    for r in recs:
        groups[r["エリアマネージャー"]].append(r)
    for idx, (am, rs) in enumerate(sorted(groups.items(), key=lambda x: (x[0] == R.AM_MISSING_LABEL, x[0]))):
        sh = wb.create_sheet(_safe_sheet(am, used))
        sh["A1"] = f"{am} 様 担当　契約更新対象者（{period_label}満了）{len(rs)}名"; sh["A1"].font = TITLE_FONT
        sh["A2"] = f"「更新可否」欄に 更新／終了／要相談 のいずれかをご記入のうえ、{jdate(deadline)}までにご回答ください。黄色セルが記入欄です。"
        _write_table(sh, 4, headers, sorted(rs, key=lambda r: (r["現契約終了日(計算)"], r["職員ID"])), name=f"T{idx:02d}")
        _autowidth(sh, from_row=4)
        col = get_column_letter(headers.index("更新可否(AM記入)") + 1)
        dv = DataValidation(type="list", formula1='"更新,終了,要相談"', allow_blank=True)
        sh.add_data_validation(dv); dv.add(f"{col}5:{col}{4 + len(rs)}")
        sh.freeze_panes = "C5"
        sh.page_setup.orientation = "landscape"; sh.page_setup.fitToWidth = 1
        sh.sheet_properties.pageSetUpPr.fitToPage = True; sh.page_setup.fitToHeight = 0

    ex = wb.create_sheet("除外(退職予定)")
    ex["A1"] = "退職予定のため対象から除外した職員（念のため確認用）"; ex["A1"].font = TITLE_FONT
    exr = [target_record(t) for t in res.excluded]
    _write_table(ex, 3, ["職員ID", "氏名", "エリア", "エリアマネージャー", "契約期間(月)", "現契約終了日(計算)"], exr, highlight=False)
    _autowidth(ex, from_row=3)
    iss = wb.create_sheet("データ要確認")
    iss["A1"] = "データ不備・要確認一覧（エラー＝計算不能のため対象判定できず）"; iss["A1"].font = TITLE_FONT
    _write_table(iss, 3, ["区分", "職員ID", "氏名", "エリア", "エリアマネージャー", "内容"], res.issues, highlight=False)
    _autowidth(iss, max_w=70, from_row=3)
    wb.save(path)


# ------------------------------------------------------------------ 依頼文
MSG = """件名：【ご確認依頼】{period}契約満了職員の更新可否について（{count}名）

{am} 様

いつもお世話になっております。人事労務担当でございます。

{period}に契約期間が満了する、{am}様ご担当の有期契約職員について、
契約更新可否のご確認をお願いしたく、ご連絡いたしました。

■ 対象者数：{count}名（うち3ヶ月契約 {n3}名）
■ 満了日：{first_end} ～ {last_end}
■ ご回答期限：{deadline}

【対象職員】
{lines}

添付の一覧（シート「{sheet}」）の「更新可否」欄に、
「更新」「終了」「要相談」のいずれかをご記入のうえ、ご返信くださいますようお願いいたします。
{notice}
契約期間の短縮・労働条件の変更・契約終了をご検討の場合は、個別に対応が必要となりますので、
お手数ですが備考欄にご記載いただくか、事前に人事労務担当までご相談ください。

ご多用のところ恐れ入りますが、何卒よろしくお願い申し上げます。

――――――――――――――――
人事労務担当　{{担当者名}}
内線：{{内線番号}}／メール：{{メールアドレス}}
――――――――――――――――
※本文は自動生成した下書きです（デモ・架空データ）。送信前に必ず内容をご確認ください。
"""


def build_messages(outdir, recs, deadline, period_label):
    os.makedirs(outdir, exist_ok=True)
    groups = defaultdict(list)
    for r in recs:
        groups[r["エリアマネージャー"]].append(r)
    paths = []
    for am, rs in sorted(groups.items()):
        rs = sorted(rs, key=lambda r: (r["現契約終了日(計算)"], r["職員ID"]))
        lines = "\n".join(
            f"・{r['職員ID']}　{r['氏名']}　{r['契約期間(月)']}ヶ月契約　満了日 {r['現契約終了日(計算)'].strftime('%Y/%m/%d')}"
            + (f"　★雇止め予告期限 {r['雇止め予告期限(30日前)'].strftime('%m/%d')}"
               if r["雇止め予告期限(30日前)"] and r["雇止め予告期限(30日前)"] <= deadline else "")
            + (f"　［{r['フラグ']}］" if r["フラグ"] else "") for r in rs)
        early = [r for r in rs if r["雇止め予告期限(30日前)"]]
        notice = ""
        if early:
            first = min(r["雇止め予告期限(30日前)"] for r in early)
            notice = (f"\nなお、{len(early)}名は更新しない場合に30日前までの雇止め予告が必要となる可能性があります"
                      f"（最も早い予告期限：{jdate(first)}）。終了をご検討の方は早めにお知らせください。\n")
            urgent = [r for r in early if r["雇止め予告期限(30日前)"] <= deadline]
            if urgent:
                notice += (f"特に★印の{len(urgent)}名は、ご回答期限より前に予告期限を迎えます"
                           f"（既に経過している場合を含みます）。終了の可能性がある場合は至急ご連絡ください。\n")
        body = MSG.format(period=period_label, count=len(rs), am=am,
                          n3=sum(1 for r in rs if r["契約期間(月)"] == 3),
                          first_end=jdate(rs[0]["現契約終了日(計算)"]), last_end=jdate(rs[-1]["現契約終了日(計算)"]),
                          deadline=jdate(deadline), lines=lines, sheet=am[:28], notice=notice)
        if am == R.AM_MISSING_LABEL:
            body = "※担当エリアマネージャー未設定のため、宛先を特定してから使用してください。\n\n" + body
        safe_am = re.sub(r'[^\w一-龥ぁ-んァ-ン]', '_', am)
        p = os.path.join(outdir, f"確認依頼_{safe_am}.txt")
        with open(p, "w", encoding="utf-8") as f:
            f.write(body)
        paths.append(p)
    return paths


# ------------------------------------------------------------------ main
def naive_comparison(rows, res: R.ExtractResult):
    correct3 = {t.row[R.COL_ID] for t in res.targets if t.months == 3}
    ids3 = {r[R.COL_ID] for r in rows if str(r.get(R.COL_PERIOD, "")).strip() == "3"}
    out = []
    for label, fn in [("旧方式A: 終了日列をそのまま参照", R.naive_end_column),
                      ("旧方式B: 3ヶ月=四半期末一斉満了と仮定", R.naive_quarter),
                      ("旧方式C: 開始月だけで周期判定", R.naive_start_month)]:
        got = set()
        for (y, m) in [R.add_months_ym(res.year, res.month, i)
                       for i in range((res.end_year * 12 + res.end_month) - (res.year * 12 + res.month) + 1)]:
            got |= fn(rows, y, m)
        got3 = got & ids3
        out.append({"方式": label, "抽出数(3ヶ月)": len(got3), "正しい対象(3ヶ月)": len(correct3),
                    "漏れ": len(correct3 - got3), "誤抽出": len(got3 - correct3)})
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default="data/staff.csv")
    ap.add_argument("--month", required=True, help="対象月 YYYY-MM")
    ap.add_argument("--lookahead", type=int, default=0, help="何ヶ月先の月末まで含めるか（1=翌月末まで）")
    ap.add_argument("--convention", choices=R.CONVENTIONS, default="civil")
    ap.add_argument("--deadline", help="AM回答期限 YYYY-MM-DD（既定: 対象月の前月20日）")
    ap.add_argument("--out", default="output")
    a = ap.parse_args(argv)
    y, m = R.parse_ym(a.month)
    deadline = date.fromisoformat(a.deadline) if a.deadline else date(*R.add_months_ym(y, m, -1), 20)
    rows = load_rows(a.input)
    res = R.extract(rows, y, m, a.lookahead, a.convention)
    tag = a.month + (f"_+{a.lookahead}" if a.lookahead else "")
    period_label = f"{y}年{m}月" if not a.lookahead else f"{y}年{m}月～{res.end_year}年{res.end_month}月"
    recs = [target_record(t) for t in res.targets]
    naive = naive_comparison(rows, res)
    os.makedirs(a.out, exist_ok=True)
    xlsx = os.path.join(a.out, f"更新対象者_{tag}.xlsx")
    build_xlsx(xlsx, res, recs, naive, deadline, period_label)
    csvp = os.path.join(a.out, f"更新対象者_{tag}.csv")
    with open(csvp, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(recs[0].keys()) if recs else ["職員ID"])
        w.writeheader(); w.writerows(recs)
    msgs = build_messages(os.path.join(a.out, f"messages_{tag}"), recs, deadline, period_label)

    print(f"対象期間: {period_label}  規約: {a.convention}  入力: {len(rows)}行")
    print(f"対象者: {len(recs)}名  （退職予定で除外 {len(res.excluded)}名 / データ要確認 {len(res.issues)}件）")
    print("契約期間別:", dict(sorted(Counter(r['契約期間(月)'] for r in recs).items())))
    print("AM別:", dict(sorted(Counter(r['エリアマネージャー'] for r in recs).items())))
    for n in naive:
        print(f"  {n['方式']}: 3ヶ月抽出 {n['抽出数(3ヶ月)']} / 正 {n['正しい対象(3ヶ月)']} → 漏れ {n['漏れ']}・誤抽出 {n['誤抽出']}")
    print("出力:", xlsx, csvp, f"依頼文 {len(msgs)}件 -> {os.path.dirname(msgs[0]) if msgs else '-'}")
    return res


if __name__ == "__main__":
    main()
