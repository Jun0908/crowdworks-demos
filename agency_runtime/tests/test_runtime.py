import copy
import json
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest

from agency_runtime.adapters import process, run_assignment
from agency_runtime.core import AgencyError, Ledger, SPEC

REPO = SPEC.parents[1]


@pytest.fixture
def ledger(tmp_path):
    result = Ledger(tmp_path / 'agency')
    result.init()
    for cap in ['article_draft', 'artifact_review', 'delivery_package', 'contract_extract', 'leads_clean', 'invoice_extract']:
        result.configure_capability(cap, 'ready', 'Local test fixture; not native Bot readiness')
    return result


def brief():
    return {'kind': 'job_brief', 'title': 'Local sample', 'objective': 'Test real artifacts and their quality gates',
            'category': 'data', 'deliverables': [{'name': 'report', 'format': 'md', 'description': 'Checked output'}],
            'inputArtifactIds': [], 'acceptanceCriteria': ['Evidence is present']}


def plan(job, capability='article_draft'):
    steps = []
    for sid, role, cap, depends in [('step_make', 'creator', capability, []),
                                   ('step_qa', 'qa', 'artifact_review', ['step_make']),
                                   ('step_delivery', 'delivery', 'delivery_package', ['step_qa'])]:
        steps.append({'id': sid, 'role': role, 'capabilityId': cap, 'dependsOn': depends,
                      'instructions': 'Read and perform the acceptance checks', 'inputArtifactIds': [],
                      'acceptanceCriteria': ['Return verified output']})
    return {'kind': 'plan', 'jobId': job['jobId'], 'summary': 'Local test', 'questions': [], 'steps': steps}


def prepared(ledger, key='one', capability='article_draft'):
    job = ledger.create(brief(), 'sample', key)
    ledger.set_plan(job['jobId'], plan(job, capability))
    return job


def result(assignment, artifacts=(), outcome='succeeded'):
    checks = [{'id': ident, 'status': 'PASS', 'evidence': 'Test fixture check'} for ident in assignment['requiredCheckIds']]
    if assignment['step']['role'] == 'qa':
        checks += [{'id': 'file_' + a['id'], 'status': 'PASS', 'evidence': 'Inspected fixture content', 'artifactId': a['id']} for a in artifacts]
    return {'kind': 'step_result', 'stepId': assignment['stepId'], 'outcome': outcome,
            'summary': 'Local test result', 'artifactIds': [a['id'] for a in artifacts], 'sourceIds': [],
            'checks': checks, 'blockers': ['Test blocker'] if outcome == 'blocked' else []}


def make(ledger, job, tmp_path):
    assignment = ledger.next_assignment(job['jobId'])
    file = tmp_path / 'report.md'
    file.write_text('# Sample\nVerified fixture content', encoding='utf-8')
    artifact = ledger.artifact(assignment['assignmentId'], file)
    ledger.finish(assignment['assignmentId'], result(assignment, [artifact]))
    return assignment, artifact


def test_request_replay_and_conflict(ledger):
    a = ledger.create(brief(), 'sample', 'stable')
    assert ledger.create(brief(), 'sample', 'stable')['jobId'] == a['jobId']
    b = brief(); b['title'] = 'Different'
    with pytest.raises(AgencyError, match='different input'): ledger.create(b, 'sample', 'stable')


@pytest.mark.parametrize('problem', ['cycle', 'missing', 'duplicate', 'noqa', 'bypass', 'question', 'foreign'])
def test_bad_plans(ledger, problem):
    job = ledger.create(brief(), 'sample', problem)
    p = plan(job)
    if problem == 'cycle': p['steps'][0]['dependsOn'] = ['step_delivery']
    if problem == 'missing': p['steps'][0]['dependsOn'] = ['step_missing']
    if problem == 'duplicate': p['steps'][1]['id'] = 'step_make'
    if problem == 'noqa': p['steps'][1]['role'] = 'creator'
    if problem == 'bypass': p['steps'][1]['dependsOn'] = []; p['steps'][2]['dependsOn'] = ['step_make']
    if problem == 'question': p['questions'] = [{'id': 'q_info', 'question': 'Missing input?', 'blocking': True}]
    if problem == 'foreign': p['jobId'] = ledger.create(brief(), 'sample', 'other')['jobId']
    with pytest.raises(AgencyError): ledger.set_plan(job['jobId'], p)


def test_concurrent_dispatch_is_idempotent(ledger):
    job = prepared(ledger)
    with ThreadPoolExecutor(max_workers=4) as pool:
        assignments = list(pool.map(lambda _: Ledger(ledger.root).next_assignment(job['jobId']), range(4)))
    assert len({a['assignmentId'] for a in assignments}) == 1
    assert ledger.show(job['jobId'])['steps']['step_make']['attempt'] == 1
    other = prepared(ledger, 'two')
    with pytest.raises(AgencyError, match='Another job'): ledger.next_assignment(other['jobId'])


def test_unconnected_capability_blocks_dispatch(ledger):
    ledger.configure_capability('article_draft', 'needs_connection', 'Missing service')
    job = prepared(ledger)
    assert ledger.next_assignment(job['jobId'])['status'] == 'waiting_connection'
    ledger.configure_capability('article_draft', 'ready', 'Actual capability check')
    ledger.control(job['jobId'], 'resume')
    assert ledger.next_assignment(job['jobId'])['stepId'] == 'step_make'


def test_pause_draft_and_active_work(ledger, tmp_path):
    job = ledger.create(brief(), 'sample', 'draft')
    ledger.control(job['jobId'], 'pause')
    assert ledger.control(job['jobId'], 'resume')['status'] == 'draft'
    ledger.set_plan(job['jobId'], plan(job))
    a = ledger.next_assignment(job['jobId'])
    ledger.control(job['jobId'], 'pause')
    file = tmp_path / 'a.md'; file.write_text('valid')
    artifact = ledger.artifact(a['assignmentId'], file)
    assert ledger.finish(a['assignmentId'], result(a, [artifact]))['status'] == 'paused'
    with pytest.raises(AgencyError): ledger.next_assignment(job['jobId'])
    assert ledger.control(job['jobId'], 'resume')['status'] == 'queued'


def test_cancel_rejects_late_result(ledger):
    job = prepared(ledger); a = ledger.next_assignment(job['jobId'])
    ledger.control(job['jobId'], 'cancel')
    with pytest.raises(AgencyError, match='Stale'): ledger.finish(a['assignmentId'], result(a, outcome='blocked'))


def test_integrity_and_result_replay(ledger, tmp_path):
    job = prepared(ledger); a, artifact = make(ledger, job, tmp_path)
    ledger.finish(a['assignmentId'], result(a, [artifact]))
    changed = result(a, [artifact]); changed['summary'] = 'Changed'
    with pytest.raises(AgencyError, match='Different result'): ledger.finish(a['assignmentId'], changed)
    Path(artifact['path']).write_text('tampered')
    with pytest.raises(AgencyError, match='changed artifact'): ledger.next_assignment(job['jobId'])


def test_qa_requires_criteria_and_all_current_files(ledger, tmp_path):
    job = prepared(ledger); _, artifact = make(ledger, job, tmp_path)
    qa = ledger.next_assignment(job['jobId'])
    with pytest.raises(AgencyError, match='inspect each'): ledger.finish(qa['assignmentId'], result(qa))
    r = result(qa, [artifact]); r['checks'] = r['checks'][1:]
    with pytest.raises(AgencyError, match='every acceptance'): ledger.finish(qa['assignmentId'], r)
    ledger.finish(qa['assignmentId'], result(qa, [artifact]))
    delivery = ledger.next_assignment(job['jobId'])
    package = ledger.package(delivery['assignmentId'])
    assert Path(package['manifestFile']).is_file()
    assert artifact['id'] in Path(package['manifestFile']).read_text(encoding='utf-8')
    output = ledger.artifact(delivery['assignmentId'], package['deliveryFile'])
    assert ledger.finish(delivery['assignmentId'], result(delivery, [output]))['status'] == 'ready_for_review'
    assert ledger.accept(job['jobId'])['status'] == 'completed'


def test_revision_invalidates_downstream_and_caps_retry(ledger, tmp_path):
    job = prepared(ledger); first, old = make(ledger, job, tmp_path)
    qa = ledger.next_assignment(job['jobId'])
    ledger.finish(qa['assignmentId'], result(qa, outcome='blocked'))
    ledger.revise(job['jobId'], 'step_make', 'Content needs revision')
    a = ledger.next_assignment(job['jobId'])
    assert a['attempt'] == 2
    assert old['id'] not in [v['id'] for v in a['inputArtifacts']]
    with pytest.raises(AgencyError, match='Stale'): ledger.finish(a['assignmentId'], result(a, [old]))
    ledger.finish(a['assignmentId'], result(a, outcome='blocked'))
    ledger.revise(job['jobId'], 'step_make', 'Second try')
    a = ledger.next_assignment(job['jobId']); ledger.finish(a['assignmentId'], result(a, outcome='blocked'))
    with pytest.raises(AgencyError, match='Revision unavailable'): ledger.revise(job['jobId'], 'step_make', 'Third try')


def test_imported_input_copied_and_cross_job_refs_rejected(ledger, tmp_path):
    job = ledger.create(brief(), 'sample', 'one')
    file = tmp_path / 'input.csv'; file.write_text('a,b\n1,2')
    artifact = ledger.import_input(job['jobId'], file)
    assert ledger.import_input(job['jobId'], file)['id'] == artifact['id']
    file.write_text('changed original')
    assert Path(artifact['path']).read_text() == 'a,b\n1,2'
    other = ledger.create(brief(), 'sample', 'two')
    p = plan(other); p['steps'][0]['inputArtifactIds'] = [artifact['id']]
    with pytest.raises(AgencyError, match='Cross-job'): ledger.set_plan(other['jobId'], p)


def test_opportunity_deduplicates_and_rejects_url(ledger):
    ledger.opportunity('stable', 'Restaurant', 'https://example.org/', 'Inspected sample')
    ledger.opportunity('stable', 'Updated', 'https://example.org/', 'New inspected sample')
    index = json.loads((ledger.root / 'opportunities/index.json').read_text())
    assert len(index) == 1 and index[0]['title'] == 'Updated'
    with pytest.raises(AgencyError): ledger.opportunity('bad', 'Bad', 'file:///private', 'Bad')


@pytest.mark.parametrize('tool', ['contract_extract', 'leads_clean', 'invoice_extract'])
def test_processing_adapters_create_real_files(ledger, tmp_path, tool):
    if tool == 'contract_extract':
        payload = {'csv_text': (REPO / 'web/dist/sample.csv').read_text(encoding='utf-8-sig'), 'month': '2026-11'}
    elif tool == 'leads_clean':
        payload = {'data': json.loads((REPO / 'web/dist/data/leads.json').read_text(encoding='utf-8')), 'sample': True}
    else:
        payload = {'documents': [{'name': 'test.pdf', 'lines': ['請求書', '株式会社テスト', '請求日 2026年10月1日', '支払期日 2026年10月31日', '合計金額 11,000円']}], 'today': '2026-10-11'}
    job = prepared(ledger, tool, tool); a = ledger.next_assignment(job['jobId'])
    output = run_assignment(ledger, a['assignmentId'], payload)
    assert ledger.show(job['jobId'])['status'] == 'running'  # no fake QA success
    assert output['summary']['externalRequests'] == 0
    from openpyxl import load_workbook
    xlsx = next(v for v in output['artifacts'] if v['name'].endswith('.xlsx'))
    assert len(load_workbook(xlsx['path']).sheetnames) >= 1
    if tool == 'contract_extract': assert len(output['summary']['records']) == 80
    if tool == 'leads_clean': assert output['summary']['summary']['final_rows'] == 14
    if tool == 'invoice_extract':
        assert output['summary']['records'][0]['total'] == 11000
        assert 'not verified' in output['summary']['records'][0]['method']
    ledger.finish(a['assignmentId'], result(a, output['artifacts']))


def test_cli_restart_and_json_exit_codes(tmp_path):
    args = [sys.executable, '-m', 'agency_runtime', '--root', str(tmp_path / 'ledger')]
    first = subprocess.run(args + ['init'], capture_output=True, text=True, encoding='utf-8')
    assert first.returncode == 0 and json.loads(first.stdout)['ok']
    after = subprocess.run(args + ['show'], capture_output=True, text=True, encoding='utf-8')
    assert json.loads(after.stdout)['data'] == []
    bad = subprocess.run(args + ['show', 'missing'], capture_output=True, text=True, encoding='utf-8')
    assert bad.returncode == 1 and not json.loads(bad.stderr)['ok']


def test_adapter_replay_reuses_files_and_rejects_changed_input(ledger, monkeypatch):
    import agency_runtime.adapters as adapters
    job = prepared(ledger, 'replay', 'contract_extract')
    a = ledger.next_assignment(job['jobId'])
    payload = {'csv_text': (REPO / 'web/dist/sample.csv').read_text(encoding='utf-8-sig'), 'month': '2026-11'}
    first = run_assignment(ledger, a['assignmentId'], payload)
    def unexpected(*args):
        raise AssertionError('Processing must not run twice')
    monkeypatch.setattr(adapters, 'process', unexpected)
    second = run_assignment(ledger, a['assignmentId'], payload)
    assert first['artifacts'] == second['artifacts']
    with pytest.raises(AgencyError, match='input changed'):
        run_assignment(ledger, a['assignmentId'], {**payload, 'month': '2026-12'})
    Path(first['artifacts'][0]['path']).write_bytes(b'changed')
    with pytest.raises(AgencyError, match='changed artifact'):
        ledger.finish(a['assignmentId'], result(a, first['artifacts']))


def test_init_preserves_unmanaged_workspace(tmp_path):
    root = tmp_path / 'previous'
    (root / 'config').mkdir(parents=True)
    file = root / 'config' / 'company-policy.json'; file.write_text('Previous Bot work')
    with pytest.raises(AgencyError, match='Existing unmanaged'):
        Ledger(root).init()
    assert file.read_text() == 'Previous Bot work'
