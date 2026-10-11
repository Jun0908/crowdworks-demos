#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
炭火焼肉店 営業リスト作成スクリプト（Google Places API (New) Text Search）

- API キーは環境変数 GOOGLE_PLACES_API_KEY から読み込みます（ハードコード禁止）。
- --mock を付けると mock/ 配下の JSON（searchText レスポンスと同じ形）を読み込み、
  API へは一切アクセスしません。

使い方:
  export GOOGLE_PLACES_API_KEY=xxxxx
  python3 yakiniku_leads.py                 # 本番（全8エリア × 既定キーワード）
  python3 yakiniku_leads.py --mock          # モック（架空データ）
  python3 yakiniku_leads.py --areas 青梅 所沢 --keywords 炭火焼肉 焼肉
  python3 yakiniku_leads.py --review-source details   # 検索は安いSKU、口コミはDetailsで補完
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import os
import random
import re
import sys
import time
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import requests

BASE_DIR = Path(__file__).resolve().parent
SEARCH_URL = "https://places.googleapis.com/v1/places:searchText"
DETAILS_URL = "https://places.googleapis.com/v1/places/{place_id}"

log = logging.getLogger("yakiniku_leads")

# ---------------------------------------------------------------------------
# エリア定義（CrowdWorks 案件 13516755 記載の 8 エリア）
#   address_keys: 住所にこの文字列が含まれていればそのエリアと判定
#   ※「入間」は「入間郡（三芳町・毛呂山町・越生町）」と区別するため「入間市」で判定
#   center: locationBias 用の概略中心座標（市役所付近の概算値）
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Area:
    name: str            # 表示名
    query_name: str      # 検索クエリに使う名称
    address_keys: tuple  # 住所判定キー
    lat: float
    lng: float
    radius_m: float = 6000.0


AREAS: list[Area] = [
    Area("青梅", "青梅市", ("東京都青梅市",), 35.7880, 139.2758, 9000),
    Area("瑞穂", "瑞穂町", ("東京都西多摩郡瑞穂町", "東京都瑞穂町"), 35.7717, 139.3540, 4000),
    Area("福生", "福生市", ("東京都福生市",), 35.7387, 139.3268, 3500),
    Area("羽村", "羽村市", ("東京都羽村市",), 35.7672, 139.3110, 3500),
    Area("昭島", "昭島市", ("東京都昭島市",), 35.7057, 139.3535, 4000),
    Area("所沢", "所沢市", ("埼玉県所沢市",), 35.7994, 139.4687, 7000),
    Area("飯能", "飯能市", ("埼玉県飯能市",), 35.8557, 139.3277, 12000),
    Area("入間", "入間市", ("埼玉県入間市",), 35.8358, 139.3911, 6000),
]
AREA_BY_NAME = {a.name: a for a in AREAS}

DEFAULT_KEYWORDS = ["炭火焼肉", "焼肉", "七輪 焼肉", "ホルモン"]

# 炭火判定キーワード
CHARCOAL_WORDS = ["炭火", "七輪", "備長炭", "炭焼", "炭焼き", "炭火焼", "木炭", "オガ炭", "七厘"]
# 口コミのみで使う弱めの語（「炭の香り」等）
REVIEW_ONLY_WORDS = ["炭の香り", "炭で焼", "炭を使", "炭の匂い", "炭火の香り"]
# 口コミで炭火否定・ガス等を示唆する語
NEGATIVE_WORDS = ["ガスロースター", "無煙ロースター", "ガス火", "炭火ではない", "炭火じゃない",
                  "炭じゃない", "ガスで焼く", "電気ロースター"]
# 焼肉関連性の判定語・タイプ
YAKINIKU_WORDS = ["焼肉", "焼き肉", "ホルモン", "七輪", "ジンギスカン", "韓国料理", "サムギョプサル",
                  "炭火焼肉", "やきにく", "ヤキニク", "BBQ", "バーベキュー"]
YAKINIKU_TYPES = {"yakiniku_restaurant", "barbecue_restaurant", "korean_restaurant",
                  "korean_barbecue_restaurant"}

BASE_FIELDS = [
    "places.id", "places.displayName", "places.formattedAddress", "places.googleMapsUri",
    "places.types", "places.primaryType", "places.businessStatus",
]
ENTERPRISE_FIELDS = ["places.nationalPhoneNumber", "places.rating", "places.userRatingCount",
                     "places.websiteUri"]
ATMOSPHERE_FIELDS = ["places.editorialSummary", "places.reviews"]

OUTPUT_COLUMNS = ["エリア", "店名", "住所", "電話", "評価", "口コミ件数", "GoogleマップURL",
                  "公式サイトURL", "炭火判定", "判定根拠", "place_id", "ヒット検索語", "備考"]
EXCLUDED_COLUMNS = ["除外理由", "店名", "住所", "電話", "place_id", "ヒット検索語", "備考"]


# ---------------------------------------------------------------------------
# API クライアント
# ---------------------------------------------------------------------------
class PlacesClient:
    """Places API (New) クライアント。リトライ・レート制御付き。"""

    def __init__(self, api_key: str, min_interval: float = 0.2, max_retries: int = 5,
                 timeout: float = 30.0):
        self.api_key = api_key
        self.min_interval = min_interval
        self.max_retries = max_retries
        self.timeout = timeout
        self.session = requests.Session()
        self._last = 0.0
        self.request_count = {"searchText": 0, "details": 0}

    def _throttle(self) -> None:
        wait = self.min_interval - (time.monotonic() - self._last)
        if wait > 0:
            time.sleep(wait)
        self._last = time.monotonic()

    def _request(self, method: str, url: str, field_mask: str, body: dict | None = None,
                 params: dict | None = None) -> dict:
        headers = {
            "Content-Type": "application/json",
            "X-Goog-Api-Key": self.api_key,
            "X-Goog-FieldMask": field_mask,
        }
        for attempt in range(1, self.max_retries + 1):
            self._throttle()
            try:
                r = self.session.request(method, url, headers=headers, json=body, params=params,
                                         timeout=self.timeout)
            except requests.RequestException as e:
                err = f"network error: {e}"
            else:
                if r.status_code == 200:
                    return r.json()
                # 4xx（429以外）はリトライしても無駄なので即エラー
                if r.status_code not in (429, 500, 502, 503, 504):
                    raise RuntimeError(f"Places API error {r.status_code}: {r.text[:500]}")
                err = f"HTTP {r.status_code}: {r.text[:200]}"
            backoff = min(60.0, (2 ** attempt) + random.uniform(0, 1))
            log.warning("retry %d/%d after %.1fs (%s)", attempt, self.max_retries, backoff, err)
            time.sleep(backoff)
        raise RuntimeError(f"Places API: retries exhausted for {url}")

    def search_text(self, body: dict, field_mask: str) -> dict:
        self.request_count["searchText"] += 1
        return self._request("POST", SEARCH_URL, field_mask, body=body)

    def place_details(self, place_id: str, field_mask: str) -> dict:
        self.request_count["details"] += 1
        return self._request("GET", DETAILS_URL.format(place_id=place_id), field_mask,
                             params={"languageCode": "ja", "regionCode": "JP"})


class MockPlacesClient:
    """mock/index.json に従い、searchText レスポンス形式の JSON を返す。"""

    def __init__(self, mock_dir: Path):
        self.mock_dir = mock_dir
        self.index = json.loads((mock_dir / "index.json").read_text(encoding="utf-8"))
        self.request_count = {"searchText": 0, "details": 0}

    def _load(self, fname: str | None) -> dict:
        if not fname:
            return {}
        return json.loads((self.mock_dir / fname).read_text(encoding="utf-8"))

    def search_text(self, body: dict, field_mask: str) -> dict:
        self.request_count["searchText"] += 1
        if body.get("pageToken"):
            return self._load(self.index.get("pages", {}).get(body["pageToken"]))
        return self._load(self.index.get("queries", {}).get(body["textQuery"]))

    def place_details(self, place_id: str, field_mask: str) -> dict:
        self.request_count["details"] += 1
        return self._load(self.index.get("details", {}).get(place_id))


# ---------------------------------------------------------------------------
# 検索
# ---------------------------------------------------------------------------
def build_field_mask(review_source: str, details: bool = False) -> str:
    fields = BASE_FIELDS + ENTERPRISE_FIELDS
    if review_source == "search" or details:
        fields = fields + ATMOSPHERE_FIELDS
    if details:  # Place Details は "places." 接頭辞なし
        return ",".join(f.replace("places.", "", 1) for f in ATMOSPHERE_FIELDS)
    return ",".join(fields + ["nextPageToken"])


def search_area(client, area: Area, keyword: str, field_mask: str, max_pages: int = 3,
                page_size: int = 20, page_wait: float = 2.0, is_mock: bool = False) -> list[dict]:
    query = f"{keyword} {area.query_name}"
    body: dict[str, Any] = {
        "textQuery": query,
        "languageCode": "ja",
        "regionCode": "JP",
        "pageSize": page_size,
        "locationBias": {"circle": {"center": {"latitude": area.lat, "longitude": area.lng},
                                    "radius": area.radius_m}},
    }
    results: list[dict] = []
    for page in range(1, max_pages + 1):
        data = client.search_text(body, field_mask)
        places = data.get("places", []) or []
        for p in places:
            p["_query"] = query
        results.extend(places)
        log.info("  [%s] %s page%d: %d件", area.name, query, page, len(places))
        token = data.get("nextPageToken")
        if not token:
            break
        body = {**body, "pageToken": token}
        if not is_mock:
            time.sleep(page_wait)  # トークン有効化待ち（念のため）
    return results


# ---------------------------------------------------------------------------
# 正規化・判定
# ---------------------------------------------------------------------------
def nfkc(s: str | None) -> str:
    return unicodedata.normalize("NFKC", s or "")


def norm_name(s: str | None) -> str:
    s = nfkc(s).lower()
    return re.sub(r"[\s\u3000・･\-ー‐－_()（）\[\]【】「」『』]", "", s)


def norm_phone(s: str | None) -> str:
    return re.sub(r"\D", "", nfkc(s))


def norm_address(s: str | None) -> str:
    s = nfkc(s)
    s = re.sub(r"^日本[、,]\s*", "", s)
    s = re.sub(r"〒?\s*\d{3}-?\d{4}\s*", "", s)
    s = re.sub(r"[\s]", "", s)
    s = s.replace("丁目", "-").replace("番地", "-").replace("番", "-").replace("号", "")
    return s.strip("-")


def detect_area(address: str | None, areas: Iterable[Area]) -> Area | None:
    a = nfkc(address).replace(" ", "")
    for area in areas:
        if any(k in a for k in area.address_keys):
            return area
    return None


def place_text(p: dict, key: str) -> str:
    v = p.get(key)
    if isinstance(v, dict):
        return v.get("text", "") or ""
    return v or ""


def review_texts(p: dict) -> list[str]:
    out = []
    for r in p.get("reviews", []) or []:
        t = (r.get("text") or {}).get("text") or (r.get("originalText") or {}).get("text") or ""
        if t:
            out.append(nfkc(t))
    return out


def find_words(text: str, words: list[str]) -> list[str]:
    return [w for w in words if w in text]


def judge_charcoal(p: dict) -> tuple[str, str]:
    """炭火判定 → (確定/可能性/不明, 根拠)"""
    name = nfkc(place_text(p, "displayName"))
    summary = nfkc(place_text(p, "editorialSummary"))
    reviews = review_texts(p)
    reasons: list[str] = []
    level = "不明"

    hit = find_words(name, CHARCOAL_WORDS)
    if hit:
        level = "確定"
        reasons.append(f"店名に「{hit[0]}」")
    hit = find_words(summary, CHARCOAL_WORDS)
    if hit:
        level = "確定"
        reasons.append(f"editorialSummaryに「{hit[0]}」")

    pos_reviews, neg_reviews = 0, 0
    pos_word = neg_word = ""
    for t in reviews:
        h = find_words(t, CHARCOAL_WORDS + REVIEW_ONLY_WORDS)
        n = find_words(t, NEGATIVE_WORDS)
        if h:
            pos_reviews += 1
            pos_word = pos_word or h[0]
        if n:
            neg_reviews += 1
            neg_word = neg_word or n[0]
    if pos_reviews:
        reasons.append(f"口コミに「{pos_word}」({pos_reviews}/{len(reviews)}件)")
        if level == "不明":
            level = "可能性"
    if neg_reviews:
        reasons.append(f"口コミに「{neg_word}」(ガス等の可能性)")
        if level == "可能性" and neg_reviews >= pos_reviews:
            level = "不明"
    if not reasons:
        reasons.append("店名・概要・口コミ(最大5件)に炭火関連語なし" if reviews or summary
                       else "炭火関連情報なし（口コミ/概要未取得）")
    return level, " / ".join(reasons)


def is_yakiniku(p: dict) -> bool:
    text = nfkc(place_text(p, "displayName") + " " + place_text(p, "editorialSummary"))
    if find_words(text, YAKINIKU_WORDS):
        return True
    types = set(p.get("types") or []) | {p.get("primaryType") or ""}
    return bool(types & YAKINIKU_TYPES)


# ---------------------------------------------------------------------------
# 重複排除
# ---------------------------------------------------------------------------
def info_score(p: dict) -> int:
    keys = ["nationalPhoneNumber", "rating", "userRatingCount", "websiteUri", "editorialSummary"]
    return sum(1 for k in keys if p.get(k)) + len(p.get("reviews") or [])


def merge(a: dict, b: dict) -> dict:
    keep, other = (a, b) if info_score(a) >= info_score(b) else (b, a)
    merged = dict(keep)
    for k, v in other.items():
        if not merged.get(k) and v:
            merged[k] = v
    merged["_queries"] = sorted(set(a.get("_queries", [])) | set(b.get("_queries", [])))
    merged["_merged_ids"] = sorted(set(a.get("_merged_ids", [a.get("id")])) |
                                   set(b.get("_merged_ids", [b.get("id")])))
    return merged


def dedupe(places: list[dict]) -> tuple[list[dict], dict]:
    stats = {"input": len(places)}
    # 1) place_id
    by_id: dict[str, dict] = {}
    for p in places:
        p.setdefault("_queries", [p.get("_query", "")])
        pid = p.get("id")
        by_id[pid] = merge(by_id[pid], p) if pid in by_id else p
    stage1 = list(by_id.values())
    stats["after_place_id"] = len(stage1)

    # 2) 正規化店名 + 電話 / 正規化店名 + 住所
    out: list[dict] = []
    key_index: dict[tuple, int] = {}
    for p in stage1:
        n = norm_name(place_text(p, "displayName"))
        keys = []
        ph = norm_phone(p.get("nationalPhoneNumber"))
        if ph:
            keys.append(("np", n, ph))
        ad = norm_address(p.get("formattedAddress"))
        if ad:
            keys.append(("na", n, ad))
        idx = next((key_index[k] for k in keys if k in key_index), None)
        if idx is None:
            out.append(p)
            idx = len(out) - 1
        else:
            out[idx] = merge(out[idx], p)
        for k in keys:
            key_index[k] = idx
    stats["after_name_phone_address"] = len(out)
    return out, stats


# ---------------------------------------------------------------------------
# メイン処理
# ---------------------------------------------------------------------------
def enrich_with_details(client, places: list[dict]) -> None:
    """店名で炭火確定できなかった店舗のみ Place Details で口コミ・概要を取得（コスト節約）。"""
    mask = build_field_mask("details", details=True)
    for p in places:
        if find_words(nfkc(place_text(p, "displayName")), CHARCOAL_WORDS):
            continue
        if p.get("reviews") or p.get("editorialSummary"):
            continue
        try:
            d = client.place_details(p["id"], mask)
        except RuntimeError as e:
            log.warning("details failed for %s: %s", p.get("id"), e)
            continue
        for k in ("reviews", "editorialSummary"):
            if d.get(k):
                p[k] = d[k]


def to_row(p: dict, area: Area, note: str) -> dict:
    level, reason = judge_charcoal(p)
    return {
        "エリア": area.name,
        "店名": place_text(p, "displayName"),
        "住所": re.sub(r"^日本、\s*", "", p.get("formattedAddress", "") or ""),
        "電話": p.get("nationalPhoneNumber", "") or "",
        "評価": p.get("rating", ""),
        "口コミ件数": p.get("userRatingCount", ""),
        "GoogleマップURL": p.get("googleMapsUri", "") or "",
        "公式サイトURL": p.get("websiteUri", "") or "",
        "炭火判定": level,
        "判定根拠": reason,
        "place_id": p.get("id", ""),
        "ヒット検索語": " | ".join(p.get("_queries", [])),
        "備考": note,
    }


def run(args) -> int:
    areas = [AREA_BY_NAME[a] for a in args.areas] if args.areas else AREAS
    if args.mock:
        client = MockPlacesClient(Path(args.mock_dir))
        note = "【架空データ】実在しない店舗です（デモ用モック）"
    else:
        key = os.environ.get("GOOGLE_PLACES_API_KEY")
        if not key:
            log.error("環境変数 GOOGLE_PLACES_API_KEY が未設定です（--mock でモック実行可）")
            return 2
        client = PlacesClient(key, min_interval=args.min_interval)
        note = ""

    mask = build_field_mask(args.review_source)
    log.info("FieldMask: %s", mask)
    raw: list[dict] = []
    for area in areas:
        for kw in args.keywords:
            raw.extend(search_area(client, area, kw, mask, max_pages=args.max_pages,
                                   is_mock=args.mock))

    uniq, stats = dedupe(raw)

    rows, excluded = [], []
    all_areas = AREAS  # 他エリアの検索で拾った店も、住所が8エリア内なら採用
    for p in uniq:
        area = detect_area(p.get("formattedAddress"), all_areas)
        reason = None
        if area is None or (args.areas and area.name not in args.areas):
            reason = "対象エリア外"
        elif p.get("businessStatus") == "CLOSED_PERMANENTLY":
            reason = "閉業"
        elif not is_yakiniku(p):
            reason = "焼肉関連性なし"
        if reason:
            excluded.append({
                "除外理由": reason, "店名": place_text(p, "displayName"),
                "住所": p.get("formattedAddress", ""), "電話": p.get("nationalPhoneNumber", ""),
                "place_id": p.get("id", ""), "ヒット検索語": " | ".join(p.get("_queries", [])),
                "備考": note,
            })
            continue
        p["_area"] = area
        rows.append(p)

    if args.review_source == "details":
        enrich_with_details(client, rows)

    out_rows = [to_row(p, p["_area"], note) for p in rows]
    order = {a.name: i for i, a in enumerate(AREAS)}
    rank = {"確定": 0, "可能性": 1, "不明": 2}
    out_rows.sort(key=lambda r: (order[r["エリア"]], rank[r["炭火判定"]], r["店名"]))

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = args.output_name or ("mock_leads" if args.mock else "yakiniku_leads")
    csv_path = out_dir / f"{stem}.csv"
    xlsx_path = out_dir / f"{stem}.xlsx"
    excl_path = out_dir / f"{stem}_excluded.csv"
    write_csv(csv_path, OUTPUT_COLUMNS, out_rows)
    write_csv(excl_path, EXCLUDED_COLUMNS, excluded)

    summary = {
        "mode": "mock" if args.mock else "live",
        "requests": client.request_count,
        "raw_results": stats["input"],
        "after_place_id_dedupe": stats["after_place_id"],
        "after_name_phone_address_dedupe": stats["after_name_phone_address"],
        "excluded": {r: sum(1 for e in excluded if e["除外理由"] == r)
                     for r in sorted({e["除外理由"] for e in excluded})},
        "final_rows": len(out_rows),
        "judgement": {k: sum(1 for r in out_rows if r["炭火判定"] == k) for k in rank},
        "by_area": {a.name: sum(1 for r in out_rows if r["エリア"] == a.name) for a in areas},
    }
    write_xlsx(xlsx_path, out_rows, excluded, summary, args.mock)
    (out_dir / f"{stem}_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    log.info("出力: %s / %s / %s", csv_path, xlsx_path, excl_path)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


# ---------------------------------------------------------------------------
# 出力
# ---------------------------------------------------------------------------
def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:  # BOM付き（Excel対策）
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path: Path, rows: list[dict], excluded: list[dict], summary: dict,
               is_mock: bool) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="3B3B3B")
    head_font = Font(bold=True, color="FFFFFF")
    fills = {"確定": PatternFill("solid", fgColor="C6EFCE"),
             "可能性": PatternFill("solid", fgColor="FFEB9C"),
             "不明": PatternFill("solid", fgColor="EDEDED")}
    widths = {"エリア": 8, "店名": 34, "住所": 40, "電話": 15, "評価": 7, "口コミ件数": 10,
              "GoogleマップURL": 22, "公式サイトURL": 22, "炭火判定": 10, "判定根拠": 48,
              "place_id": 20, "ヒット検索語": 30, "備考": 30, "除外理由": 14}

    def sheet(ws, columns, data):
        ws.append(columns)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center")
        for r in data:
            ws.append([r.get(c, "") for c in columns])
        for i, col in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 15)
        ws.freeze_panes = "A2"
        if data:
            ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=columns[c.column - 1]
                                        in ("判定根拠", "住所", "ヒット検索語"))
                col = columns[c.column - 1]
                if col in ("GoogleマップURL", "公式サイトURL") and c.value:
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
                if col == "炭火判定" and c.value in fills:
                    c.fill = fills[c.value]

    ws = wb.active
    ws.title = "リスト"
    sheet(ws, OUTPUT_COLUMNS, rows)
    sheet(wb.create_sheet("除外"), EXCLUDED_COLUMNS, excluded)

    ws3 = wb.create_sheet("集計")
    ws3.column_dimensions["A"].width = 36
    ws3.column_dimensions["B"].width = 40
    def kv(k, v):
        ws3.append([k, v])
    kv("モード", summary["mode"])
    kv("API リクエスト数(searchText)", summary["requests"]["searchText"])
    kv("API リクエスト数(details)", summary["requests"]["details"])
    kv("検索結果 延べ件数", summary["raw_results"])
    kv("place_id 重複排除後", summary["after_place_id_dedupe"])
    kv("店名+電話/住所 重複排除後", summary["after_name_phone_address_dedupe"])
    for k, v in summary["excluded"].items():
        kv(f"除外: {k}", v)
    kv("最終件数", summary["final_rows"])
    for k, v in summary["judgement"].items():
        kv(f"炭火判定: {k}", v)
    for k, v in summary["by_area"].items():
        kv(f"エリア: {k}", v)
    for c in ws3["A"]:
        c.font = Font(bold=True)

    ws4 = wb.create_sheet("注記")
    ws4.column_dimensions["A"].width = 120
    notes = [
        "炭火判定は Google Places に「炭火」属性が無いため、店名・editorialSummary・口コミ(最大5件)の語句から推定しています。",
        "「確定」= 店名または概要に炭火関連語 / 「可能性」= 口コミに炭火関連語 / 「不明」= 根拠なし。営業前に電話等で確認してください。",
        "データ出典: Google Maps Platform（Places API）。利用・保存・表示は Google Maps Platform 利用規約に従ってください。",
    ]
    if is_mock:
        notes.insert(0, "【重要】このファイルはモック（架空）データです。店名・住所・電話番号・URL はすべて架空で、実在しません。")
    for n in notes:
        ws4.append([n])
    ws4["A1"].font = Font(bold=True, color="C00000" if is_mock else "000000")
    wb.save(path)


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="炭火焼肉店 営業リスト作成 (Places API New)")
    ap.add_argument("--mock", action="store_true", help="mock/ の JSON を使い API にアクセスしない")
    ap.add_argument("--mock-dir", default=str(BASE_DIR / "mock"))
    ap.add_argument("--out-dir", default=str(BASE_DIR / "output"))
    ap.add_argument("--output-name", default=None, help="出力ファイル名の stem")
    ap.add_argument("--areas", nargs="*", choices=list(AREA_BY_NAME), default=None)
    ap.add_argument("--keywords", nargs="*", default=DEFAULT_KEYWORDS)
    ap.add_argument("--max-pages", type=int, default=3, help="1クエリあたり最大ページ(20件/頁, 上限60件)")
    ap.add_argument("--min-interval", type=float, default=0.2, help="リクエスト間隔(秒)")
    ap.add_argument("--review-source", choices=["search", "details", "none"], default="search",
                    help="search=検索時に口コミ/概要取得(Atmosphere SKU) / details=未確定店のみDetailsで取得 / none=取得しない")
    ap.add_argument("-v", "--verbose", action="store_true")
    return ap.parse_args(argv)


if __name__ == "__main__":
    a = parse_args()
    logging.basicConfig(level=logging.INFO if a.verbose else logging.WARNING,
                        format="%(levelname)s %(message)s")
    sys.exit(run(a))
